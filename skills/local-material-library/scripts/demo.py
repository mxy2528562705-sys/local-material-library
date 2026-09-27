"""Generate fictional, redistributable fixtures. Never reads personal documents."""
import argparse
from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document


def make_pdf(path, texts):
    # Tiny fixture writer; the application uses PDFium for real parsing/rendering.
    objects = [b'', b'']
    objects.append(b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>')
    page_ids = []
    for text in texts:
        escaped = text.replace('\\', '\\\\').replace('(', '\\(').replace(')', '\\)')
        stream = f'BT /F1 20 Tf 50 700 Td ({escaped}) Tj ET'.encode('ascii')
        page_id = len(objects) + 1
        page_ids.append(page_id)
        objects.append(f'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Resources << /Font << /F1 3 0 R >> >> /Contents {page_id+1} 0 R >>'.encode())
        objects.append(b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream')
    objects[0] = b'<< /Type /Catalog /Pages 2 0 R >>'
    objects[1] = f'<< /Type /Pages /Count {len(page_ids)} /Kids [{" ".join(f"{i} 0 R" for i in page_ids)}] >>'.encode()
    output = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f'{number} 0 obj\n'.encode() + obj + b'\nendobj\n')
    xref = len(output)
    output.extend(f'xref\n0 {len(objects)+1}\n0000000000 65535 f \n'.encode())
    for pos in offsets[1:]:
        output.extend(f'{pos:010d} 00000 n \n'.encode())
    output.extend(f'trailer\n<< /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n'.encode())
    path.write_bytes(output)


def generate(output):
    root = Path(output)
    if root.exists() and any(root.iterdir()):
        raise ValueError('Demo output must be empty')
    for name in ['研究文献', '田野笔记', '图像素材']:
        (root / name).mkdir(parents=True, exist_ok=True)
    make_pdf(root / '研究文献/Regional survey.pdf', [
        'Regional survey - original layout, page one',
        'Irrigation and village records - interior page two',
        'Late-page discovery: WATERMARKET research evidence'])
    (root / '研究文献/Regional survey copy.pdf').write_bytes((root / '研究文献/Regional survey.pdf').read_bytes())
    (root / '田野笔记/水利访谈.md').write_text('# 水利与村落\n\n这是一份虚构的演示记录。关键词：水利、访谈、村落。\n', encoding='utf-8')
    doc = Document()
    doc.add_heading('资料整理清单', 0)
    doc.add_paragraph('虚构示例：整理村落史料、核对地图年代、记录资料出处。')
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = '专题'
    table.cell(0, 1).text = '地方文献'
    doc.save(root / '田野笔记/整理清单.docx')
    image = Image.new('RGB', (900, 600), '#f0f5f1')
    draw = ImageDraw.Draw(image)
    draw.line([(40, 540), (290, 390), (370, 230), (810, 70)], fill='#459aac', width=18)
    for x, y, name in [(180, 420, 'Village A'), (460, 260, 'Village B'), (700, 160, 'Village C')]:
        draw.rectangle((x, y, x+70, y+40), fill='#42715c')
        draw.text((x, y-25), name, fill='#23332c')
    draw.text((30, 30), 'FICTIONAL FIELD MAP / DEMO', fill='#23332c')
    image.save(root / '图像素材/村落示意图.png')
    return root


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', required=True)
    print(generate(parser.parse_args().output))
