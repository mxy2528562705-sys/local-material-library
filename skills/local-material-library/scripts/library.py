"""Read-only source archive, incremental SQLite index and loopback web UI."""
import argparse
from contextlib import closing, contextmanager
import hashlib
import io
import json
import mimetypes
import os
from pathlib import Path
import secrets
import shutil
import sqlite3
import subprocess
import tempfile
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, urlsplit, quote

import pypdfium2 as pdfium
from PIL import Image, ImageOps
from docx import Document

STATIC = Path(__file__).resolve().parents[1] / "assets/ui"
IMAGES = {".jpg", ".jpeg", ".png", ".webp", ".tif", ".tiff", ".bmp"}
TEXT = {".txt", ".md", ".csv"}
SUPPORTED = IMAGES | TEXT | {".pdf", ".docx"}
PDF_LOCK = threading.RLock()


def compact(text):
    return "".join(text.split()).casefold()


def file_hash(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


class Library:
    def __init__(self, source, data, rebind=False):
        self.source = Path(source).expanduser().resolve()
        self.data = Path(data).expanduser().resolve()
        if not self.source.is_dir():
            raise ValueError("资料文件夹不存在")
        if self.data == self.source or self.data in self.source.parents:
            raise ValueError("数据文件夹不能等于或包含资料文件夹")
        self.data.mkdir(parents=True, exist_ok=True)
        self.db = self.data / "library.sqlite"
        self.job_lock = threading.Lock()
        self.job = {"running": False, "message": "就绪", "error": ""}
        with self.connect() as con:
            con.executescript("""
            CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
            CREATE TABLE IF NOT EXISTS files(
              id INTEGER PRIMARY KEY AUTOINCREMENT, path TEXT UNIQUE NOT NULL,
              name TEXT NOT NULL, ext TEXT NOT NULL, size INTEGER, sha TEXT,
              category TEXT NOT NULL, tags TEXT DEFAULT '', notes TEXT DEFAULT '',
              hidden INTEGER DEFAULT 0, missing INTEGER DEFAULT 0,
              pages INTEGER DEFAULT 0, error TEXT DEFAULT '');
            CREATE TABLE IF NOT EXISTS pages(
              file_id INTEGER REFERENCES files(id) ON DELETE CASCADE,
              number INTEGER, text TEXT, status TEXT, error TEXT DEFAULT '',
              PRIMARY KEY(file_id,number));
            CREATE INDEX IF NOT EXISTS files_sha ON files(sha);
            """)
            old = con.execute("SELECT value FROM settings WHERE key='source'").fetchone()
            if old and old[0] != str(self.source) and not rebind:
                raise ValueError("数据属于另一个资料目录；迁移请使用 rebind 命令")
            con.execute("INSERT OR REPLACE INTO settings VALUES('source',?)", (str(self.source),))

    @contextmanager
    def connect(self):
        con = sqlite3.connect(self.db, timeout=30)
        con.row_factory = sqlite3.Row
        con.execute("PRAGMA foreign_keys=ON")
        con.execute("PRAGMA journal_mode=WAL")
        try:
            with con:
                yield con
        finally:
            con.close()

    def source_path(self, record):
        path = (self.source / record["path"]).resolve()
        if not path.is_relative_to(self.source) or not path.is_file():
            raise FileNotFoundError("原文件已移动或不可访问，请刷新索引")
        if path.suffix.lower() not in SUPPORTED:
            raise ValueError("不支持此文件格式")
        return path

    def record(self, file_id):
        with self.connect() as con:
            row = con.execute("SELECT * FROM files WHERE id=?", (file_id,)).fetchone()
        if not row:
            raise FileNotFoundError("条目不存在")
        return row

    def scan(self):
        paths, errors, skipped = [], [], 0
        def walk_error(error):
            errors.append(str(error))
        for folder, dirs, names in os.walk(self.source, onerror=walk_error, followlinks=False):
            dirs[:] = sorted(d for d in dirs if not d.startswith('.')
                             and not (Path(folder) / d).is_symlink()
                             and (Path(folder) / d).resolve() != self.data)
            for name in sorted(names):
                path = Path(folder) / name
                if name.startswith('.') or path.is_symlink():
                    continue
                if path.suffix.lower() in SUPPORTED:
                    paths.append(path)
                else:
                    skipped += 1
        # A failed directory enumeration must not mark unseen originals missing.
        if errors:
            raise OSError("; ".join(errors))
        relative_paths = {p.relative_to(self.source).as_posix() for p in paths}
        with self.connect() as con:
            for row in con.execute("SELECT id,path FROM files").fetchall():
                con.execute("UPDATE files SET missing=? WHERE id=?",
                            (int(row['path'] not in relative_paths), row['id']))
        for index, path in enumerate(paths, 1):
            self.job["message"] = f"导入 {index}/{len(paths)}：{path.name}"
            relative = path.relative_to(self.source).as_posix()
            try:
                sha = file_hash(path)
                size = path.stat().st_size
                with self.connect() as con:
                    row = con.execute("SELECT * FROM files WHERE path=?", (relative,)).fetchone()
                    if not row:
                        moved = con.execute("SELECT * FROM files WHERE missing=1 AND sha=?", (sha,)).fetchall()
                        if len(moved) == 1:
                            row = moved[0]
                            con.execute("UPDATE files SET path=?, name=?, ext=?, missing=0 WHERE id=?",
                                        (relative, path.name, path.suffix.lower(), row['id']))
                        else:
                            cur = con.execute("INSERT INTO files(path,name,ext,size,category) VALUES(?,?,?,?,?)",
                                              (relative, path.name, path.suffix.lower(), size,
                                               str(Path(relative).parent) if '/' in relative else '未分类'))
                            row = con.execute("SELECT * FROM files WHERE id=?", (cur.lastrowid,)).fetchone()
                    file_id = row['id']
                    if row['sha'] == sha and not row['error']:
                        continue
                    con.execute("DELETE FROM pages WHERE file_id=?", (file_id,))
                    con.execute("UPDATE files SET sha=?,size=?,missing=0,pages=0,error='提取未完成' WHERE id=?",
                                (sha, size, file_id))
                try:
                    self.extract(file_id, path)
                    if file_hash(path) != sha:
                        raise OSError("文件在导入期间发生变化，请重新刷新")
                    with self.connect() as con:
                        con.execute("UPDATE files SET error='' WHERE id=?", (file_id,))
                except Exception as error:
                    with self.connect() as con:
                        con.execute("UPDATE files SET error=? WHERE id=?", (str(error), file_id))
                    errors.append(f"{path.name}: {error}")
            except OSError as error:
                errors.append(f"{path.name}: {error}")
        self.job['message'] = f"导入完成：{len(paths)} 个文件，跳过 {skipped} 个不支持的文件，{len(errors)} 个错误"
        self.job['error'] = '\n'.join(errors)
        return {"files": len(paths), "skipped": skipped, "errors": errors}

    def put_page(self, file_id, number, text, status, error=''):
        with self.connect() as con:
            con.execute("INSERT OR REPLACE INTO pages VALUES(?,?,?,?,?)", (file_id, number, text, status, error))

    def extract(self, file_id, path):
        ext = path.suffix.lower()
        if ext == '.pdf':
            with PDF_LOCK, pdfium.PdfDocument(path) as pdf:
                with self.connect() as con:
                    con.execute("UPDATE files SET pages=? WHERE id=?", (len(pdf), file_id))
                for number in range(len(pdf)):
                    with closing(pdf[number]) as page, closing(page.get_textpage()) as textpage:
                        text = textpage.get_text_bounded()
                    self.put_page(file_id, number + 1, text, 'pending' if len(compact(text)) < 20 else 'native')
        else:
            if ext == '.docx':
                doc = Document(path)
                text = '\n'.join(p.text for p in doc.paragraphs)
                text += '\n' + '\n'.join('\t'.join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
            elif ext in TEXT:
                raw = path.read_bytes()
                try:
                    text = raw.decode('utf-8-sig')
                except UnicodeDecodeError:
                    text = raw.decode('gb18030')
            else:
                with Image.open(path) as image:
                    image.verify()
                text = ''
            with self.connect() as con:
                con.execute("UPDATE files SET pages=1 WHERE id=?", (file_id,))
            self.put_page(file_id, 1, text, 'pending' if ext in IMAGES else 'native')

    def render(self, file_id, number, width=1400):
        row = self.record(file_id)
        if number < 1 or number > row['pages']:
            raise ValueError("页码超出范围")
        path = self.source_path(row)
        if row['ext'] == '.pdf':
            with PDF_LOCK, pdfium.PdfDocument(path) as pdf, closing(pdf[number - 1]) as page:
                scale = min(width / page.get_width(), 3, (12_000_000 / (page.get_width() * page.get_height())) ** .5)
                with closing(page.render(scale=scale)) as bitmap:
                    image = bitmap.to_pil().copy()
        elif row['ext'] in IMAGES:
            with Image.open(path) as original:
                image = ImageOps.exif_transpose(original).convert('RGB')
                image.thumbnail((width, width * 3))
        else:
            raise ValueError("此格式没有原版图片预览")
        out = io.BytesIO()
        image.save(out, format='PNG')
        image.close()
        return out.getvalue()

    def ocr(self, batch=50, lang='chi_sim+eng', file_id=None, all_pages=False, retry=False):
        exe = shutil.which('tesseract')
        if not exe:
            raise ValueError("未安装 Tesseract；请按工作流说明安装 OCR 引擎和语言包")
        check = subprocess.run([exe, '--list-langs'], capture_output=True, text=True, check=True)
        installed = set(check.stdout.splitlines())
        if not set(lang.split('+')) <= installed:
            raise ValueError(f"缺少 OCR 语言包：{lang}")
        statuses = ['pending'] + (['failed'] if retry else []) + (['native'] if all_pages else [])
        where = 'p.status IN (' + ','.join('?' for _ in statuses) + ') AND f.missing=0 AND f.hidden=0 AND f.error=\'\''
        params = list(statuses)
        if all_pages and file_id is None:
            raise ValueError("全页 OCR 必须指定 --id")
        if file_id is not None:
            where += ' AND f.id=?'
            params.append(file_id)
        with self.connect() as con:
            pages = con.execute(f"SELECT p.*,f.ext FROM pages p JOIN files f ON f.id=p.file_id WHERE {where} ORDER BY f.id,p.number LIMIT ?",
                                (*params, batch)).fetchall()
        count, failed = 0, 0
        for row in pages:
            if row['ext'] not in IMAGES | {'.pdf'}:
                continue
            self.job['message'] = f"OCR {count + failed + 1}/{len(pages)}，文件 {row['file_id']} 第 {row['number']} 页"
            try:
                with tempfile.TemporaryDirectory() as temp:
                    image = Path(temp) / 'page.png'
                    image.write_bytes(self.render(row['file_id'], row['number'], 2200))
                    result = subprocess.run([exe, str(image), 'stdout', '-l', lang], capture_output=True,
                                            text=True, encoding='utf-8', timeout=180, check=True)
                self.put_page(row['file_id'], row['number'], result.stdout, 'ocr')
                count += 1
            except Exception as error:
                self.put_page(row['file_id'], row['number'], row['text'], 'failed', str(error))
                failed += 1
        self.job['message'] = f"本批 OCR 完成 {count} 页，失败 {failed} 页"
        return {"completed": count, "failed": failed}

    def listing(self, q='', category='', state='active', offset=0, limit=30):
        where, params = ["1=1"], []
        if state == 'hidden':
            where.append('f.hidden=1')
        else:
            where.append('f.hidden=0')
        if state == 'missing':
            where.append('f.missing=1')
        if state == 'duplicates':
            where.append("f.sha IN (SELECT sha FROM files WHERE missing=0 GROUP BY sha HAVING count(*)>1)")
        if state == 'pending':
            where.append("EXISTS(SELECT 1 FROM pages p WHERE p.file_id=f.id AND p.status IN ('pending','failed'))")
        if category:
            where.append('f.category=?')
            params.append(category)
        # Escape LIKE metacharacters: user input is a literal substring, including % and _.
        pattern = '%' + q.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_') + '%'
        if q:
            where.append("((f.name||' '||f.path||' '||f.category||' '||f.tags||' '||f.notes) LIKE ? ESCAPE '\\' OR EXISTS(SELECT 1 FROM pages p WHERE p.file_id=f.id AND p.text LIKE ? ESCAPE '\\'))")
            params += [pattern, pattern]
        clause = ' AND '.join(where)
        with self.connect() as con:
            total = con.execute(f'SELECT count(*) FROM files f WHERE {clause}', params).fetchone()[0]
            rows = con.execute(f'SELECT f.* FROM files f WHERE {clause} ORDER BY f.name,f.id LIMIT ? OFFSET ?', (*params, limit, offset)).fetchall()
            items = []
            for row in rows:
                item = dict(row)
                counts = dict(con.execute('SELECT status,count(*) FROM pages WHERE file_id=? GROUP BY status', (row['id'],)).fetchall())
                item['counts'] = counts
                hit = con.execute("SELECT number,text FROM pages WHERE file_id=? AND text LIKE ? ESCAPE '\\' ORDER BY number LIMIT 1",
                                  (row['id'], pattern)).fetchone() if q else None
                item['hit'] = hit['number'] if hit else 1
                text = hit['text'] if hit else ''
                start = max(0, text.casefold().find(q.casefold()) - 55) if q else 0
                item['snippet'] = text[start:start + 200]
                items.append(item)
            categories = [r[0] for r in con.execute('SELECT DISTINCT category FROM files ORDER BY category')]
        return {'items': items, 'total': total, 'categories': categories}

    def detail(self, file_id, number, q='', offset=0):
        item = dict(self.record(file_id))
        if number < 1 or number > max(1, item['pages']):
            raise ValueError('页码超出范围')
        pattern = '%' + q.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_') + '%'
        with self.connect() as con:
            row = con.execute('SELECT * FROM pages WHERE file_id=? AND number=?', (file_id, number)).fetchone()
            item['page'] = dict(row) if row else {'text': '', 'status': 'failed', 'error': item['error']}
            item['hits'] = [r[0] for r in con.execute("SELECT number FROM pages WHERE file_id=? AND text LIKE ? ESCAPE '\\' ORDER BY number LIMIT 50 OFFSET ?", (file_id, pattern, offset))] if q else []
            item['hit_total'] = con.execute("SELECT count(*) FROM pages WHERE file_id=? AND text LIKE ? ESCAPE '\\'", (file_id, pattern)).fetchone()[0] if q else 0
        return item

    def edit(self, file_id, data):
        self.record(file_id)
        allowed = {'category': 200, 'tags': 1000, 'notes': 10000, 'hidden': 1}
        with self.connect() as con:
            for key, value in data.items():
                if key not in allowed:
                    raise ValueError('未知字段')
                if key == 'hidden':
                    if type(value) is not bool:
                        raise ValueError('隐藏状态必须为布尔值')
                elif not isinstance(value, str) or len(value) > allowed[key]:
                    raise ValueError('字段过长或格式错误')
                con.execute(f'UPDATE files SET {key}=? WHERE id=?', (value, file_id))

    def start_job(self, action):
        if not self.job_lock.acquire(blocking=False):
            raise ValueError('已有导入或 OCR 任务正在运行')
        self.job = {'running': True, 'message': '准备中', 'error': ''}
        def run():
            try:
                action()
            except Exception as error:
                self.job['error'] = str(error)
                self.job['message'] = '任务失败'
            finally:
                self.job['running'] = False
                self.job_lock.release()
        threading.Thread(target=run, daemon=True).start()


def handler_for(library, title):
    token = secrets.token_urlsafe(32)
    class Handler(BaseHTTPRequestHandler):
        def log_message(self, *args):
            pass

        def send(self, content, mime='application/json; charset=utf-8', status=200, extra=None):
            if not isinstance(content, bytes):
                content = json.dumps(content, ensure_ascii=False).encode()
            self.send_response(status)
            self.send_header('Content-Type', mime)
            self.send_header('Content-Length', str(len(content)))
            self.send_header('X-Content-Type-Options', 'nosniff')
            self.send_header('Referrer-Policy', 'no-referrer')
            self.send_header('Cache-Control', 'no-store')
            self.send_header('Content-Security-Policy', "default-src 'self'; script-src 'self'; style-src 'self'; img-src 'self' data:; object-src 'none'; frame-ancestors 'none'; base-uri 'none'")
            for key, value in (extra or {}).items():
                self.send_header(key, value)
            self.end_headers()
            try:
                self.wfile.write(content)
            except (BrokenPipeError, ConnectionResetError):
                pass

        def trusted(self):
            expected = f'127.0.0.1:{self.server.server_port}'
            return self.headers.get('Host') == expected

        def do_GET(self):
            if not self.trusted():
                return self.send({'error': 'Invalid Host'}, status=403)
            parsed = urlsplit(self.path)
            query = {k: v[0] for k, v in parse_qs(parsed.query).items()}
            try:
                if parsed.path in {'/', '/app.js', '/styles.css'}:
                    name = 'index.html' if parsed.path == '/' else parsed.path[1:]
                    return self.send((STATIC / name).read_bytes(), mimetypes.guess_type(name)[0] + '; charset=utf-8')
                if parsed.path == '/api/status':
                    return self.send({'title': title, 'job': library.job, 'token': token, 'source': str(library.source)})
                if parsed.path == '/api/files':
                    return self.send(library.listing(query.get('q', '')[:500], query.get('category', ''), query.get('state', 'active'), max(0, int(query.get('offset', 0)))))
                parts = parsed.path.strip('/').split('/')
                if len(parts) == 3 and parts[:2] == ['api', 'file']:
                    file_id = int(parts[2])
                    return self.send(library.detail(file_id, int(query.get('page', 1)), query.get('q', '')[:500], max(0, int(query.get('offset', 0)))))
                if len(parts) == 4 and parts[:2] == ['api', 'page']:
                    return self.send(library.render(int(parts[2]), int(parts[3])), 'image/png')
                if len(parts) == 3 and parts[:2] == ['api', 'original']:
                    row = library.record(int(parts[2]))
                    path = library.source_path(row)
                    # Stream large originals instead of reading a whole book into memory.
                    with path.open('rb') as stream:
                        self.send_response(200)
                        self.send_header('Content-Type', 'application/octet-stream')
                        self.send_header('X-Content-Type-Options', 'nosniff')
                        self.send_header('Content-Length', str(path.stat().st_size))
                        self.send_header('Content-Disposition', "attachment; filename*=UTF-8''" + quote(row['name']))
                        self.end_headers()
                        shutil.copyfileobj(stream, self.wfile, 128 * 1024)
                    return
                self.send({'error': '不存在'}, status=404)
            except FileNotFoundError as error:
                self.send({'error': str(error)}, status=404)
            except (ValueError, OSError, sqlite3.Error) as error:
                self.send({'error': str(error)}, status=400)
            except Exception:
                self.send({'error': '无法读取此文件，请检查格式和访问权限'}, status=500)

        def do_POST(self):
            if not self.trusted() or self.headers.get('X-Library-Token') != token or self.headers.get('Origin', f'http://127.0.0.1:{self.server.server_port}') != f'http://127.0.0.1:{self.server.server_port}':
                return self.send({'error': '请求来源无效'}, status=403)
            try:
                length = int(self.headers.get('Content-Length', '0'))
                if not 0 <= length <= 32768:
                    raise ValueError('请求过大')
                data = json.loads(self.rfile.read(length) or b'{}')
                if self.path == '/api/scan':
                    library.start_job(library.scan)
                elif self.path == '/api/ocr':
                    library.start_job(lambda: library.ocr(retry=True))
                elif self.path.startswith('/api/edit/'):
                    if not isinstance(data, dict):
                        raise ValueError('无效字段')
                    library.edit(int(self.path.rsplit('/', 1)[1]), data)
                else:
                    return self.send({'error': '不存在'}, status=404)
                self.send({'ok': True})
            except (ValueError, TypeError) as error:
                self.send({'error': str(error)}, status=400)
    return Handler


def export_data(data, output):
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=True)
    db = Path(data).expanduser().resolve() / 'library.sqlite'
    if output == db.parent or (output / 'library.sqlite').exists():
        raise ValueError('备份目标必须是尚无数据库的新目录')
    with closing(sqlite3.connect(f'{db.as_uri()}?mode=ro', uri=True)) as source, closing(sqlite3.connect(output / 'library.sqlite')) as dest:
        source.backup(dest)
        source.row_factory = sqlite3.Row
        rows = [dict(row) for row in source.execute('SELECT * FROM files')]
    (output / 'metadata.json').write_text(json.dumps(rows, ensure_ascii=False, indent=2), encoding='utf-8')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('command', choices=['serve', 'scan', 'ocr', 'export', 'rebind'])
    parser.add_argument('--source')
    parser.add_argument('--data', required=True)
    parser.add_argument('--title', default='本地资料库')
    parser.add_argument('--port', type=int, default=8876)
    parser.add_argument('--batch', type=int, default=50)
    parser.add_argument('--lang', default='chi_sim+eng')
    parser.add_argument('--id', type=int)
    parser.add_argument('--all-pages', action='store_true')
    parser.add_argument('--retry', action='store_true')
    parser.add_argument('--output')
    args = parser.parse_args()
    if args.command == 'export':
        if not args.output:
            parser.error('--output is required')
        export_data(args.data, args.output)
        return
    if not args.source:
        parser.error('--source is required')
    library = Library(args.source, args.data, args.command == 'rebind')
    if args.command == 'scan':
        print(json.dumps(library.scan(), ensure_ascii=False))
    elif args.command == 'ocr':
        if args.batch < 1:
            parser.error('--batch must be positive')
        print(library.ocr(args.batch, args.lang, args.id, args.all_pages, args.retry))
    elif args.command == 'serve':
        handler = handler_for(library, args.title)
        try:
            server = ThreadingHTTPServer(('127.0.0.1', args.port), handler)
        except OSError:
            server = ThreadingHTTPServer(('127.0.0.1', 0), handler)
        with library.connect() as con:
            empty = con.execute('SELECT count(*) FROM files').fetchone()[0] == 0
        if empty:
            library.start_job(library.scan)
        print(f'http://127.0.0.1:{server.server_port}/', flush=True)
        try:
            server.serve_forever()
        except KeyboardInterrupt:
            pass
        finally:
            server.server_close()


if __name__ == '__main__':
    main()
